import os
import json
from datetime import datetime
from flask import Flask, render_template, request, jsonify, send_file
from models import db, Student, ScoreRecord
try:
    from openai import OpenAI
except ImportError:
    OpenAI = None

app = Flask(__name__)
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///database.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db.init_app(app)

# AI client (you can configure your own API key)
AI_API_KEY = os.environ.get('OPENAI_API_KEY', '')
AI_BASE_URL = os.environ.get('OPENAI_BASE_URL', 'https://api.openai.com/v1')
AI_MODEL = os.environ.get('AI_MODEL', 'gpt-3.5-turbo')

def get_ai_client():
    if AI_API_KEY and OpenAI:
        return OpenAI(api_key=AI_API_KEY, base_url=AI_BASE_URL)
    return None

# ==================== Routes ====================

@app.route('/')
def index():
    """Dashboard"""
    students = Student.query.all()
    total_students = len(students)
    
    # Calculate average scores
    total_score = sum(s.total_score for s in students)
    avg_score = round(total_score / total_students, 1) if total_students > 0 else 0
    
    # Top 5 students
    top_students = sorted(students, key=lambda x: x.total_score, reverse=True)[:5]
    
    # Recent score records
    recent_records = ScoreRecord.query.order_by(ScoreRecord.created_at.desc()).limit(10).all()
    
    return render_template('index.html',
                         total_students=total_students,
                         avg_score=avg_score,
                         top_students=top_students,
                         recent_records=recent_records)

@app.route('/students')
def students_page():
    """Student management page"""
    students = Student.query.all()
    return render_template('students.html', students=students)

@app.route('/scoring')
def scoring_page():
    """Scoring page"""
    students = Student.query.all()
    return render_template('scoring.html', students=students)

@app.route('/rankings')
def rankings_page():
    """Rankings page"""
    students = Student.query.all()
    students_sorted = sorted(students, key=lambda x: x.total_score, reverse=True)
    return render_template('rankings.html', students=students_sorted)

@app.route('/reports')
def reports_page():
    """Reports page"""
    students = Student.query.all()
    return render_template('reports.html', students=students)

# ==================== API Routes ====================

# Student APIs
@app.route('/api/students', methods=['GET'])
def get_students():
    students = Student.query.all()
    return jsonify([s.to_dict() for s in students])

@app.route('/api/students', methods=['POST'])
def add_student():
    data = request.json
    student = Student(
        name=data['name'],
        student_id=data.get('student_id', ''),
        gender=data.get('gender', '未知'),
        group=data.get('group', ''),
        phone=data.get('phone', ''),
        notes=data.get('notes', '')
    )
    db.session.add(student)
    db.session.commit()
    return jsonify(student.to_dict()), 201

@app.route('/api/students/<int:id>', methods=['PUT'])
def update_student(id):
    student = Student.query.get_or_404(id)
    data = request.json
    student.name = data.get('name', student.name)
    student.student_id = data.get('student_id', student.student_id)
    student.gender = data.get('gender', student.gender)
    student.group = data.get('group', student.group)
    student.phone = data.get('phone', student.phone)
    student.notes = data.get('notes', student.notes)
    db.session.commit()
    return jsonify(student.to_dict())

@app.route('/api/students/<int:id>', methods=['DELETE'])
def delete_student(id):
    student = Student.query.get_or_404(id)
    ScoreRecord.query.filter_by(student_id=id).delete()
    db.session.delete(student)
    db.session.commit()
    return jsonify({'message': '删除成功'})

# Score APIs
@app.route('/api/scores', methods=['POST'])
def add_score():
    data = request.json
    record = ScoreRecord(
        student_id=data['student_id'],
        score_type=data['score_type'],
        category=data.get('category', ''),
        score=data['score'],
        reason=data.get('reason', ''),
        operator=data.get('operator', '系统')
    )
    db.session.add(record)
    
    # Update student score
    student = Student.query.get(data['student_id'])
    student.total_score += data['score']
    
    db.session.commit()
    return jsonify(record.to_dict()), 201

@app.route('/api/scores/<int:student_id>', methods=['GET'])
def get_student_scores(student_id):
    records = ScoreRecord.query.filter_by(student_id=student_id)\
        .order_by(ScoreRecord.created_at.desc()).all()
    return jsonify([r.to_dict() for r in records])

@app.route('/api/scores/<int:record_id>', methods=['DELETE'])
def delete_score(record_id):
    record = ScoreRecord.query.get_or_404(record_id)
    student = Student.query.get(record.student_id)
    student.total_score -= record.score
    db.session.delete(record)
    db.session.commit()
    return jsonify({'message': '删除成功'})

# Statistics API
@app.route('/api/statistics', methods=['GET'])
def get_statistics():
    students = Student.query.all()
    if not students:
        return jsonify({'avg_score': 0, 'max_score': 0, 'min_score': 0, 
                       'score_distribution': {}, 'category_stats': {}})
    
    scores = [s.total_score for s in students]
    
    # Score distribution
    distribution = {'优秀(80-100)': 0, '良好(60-79)': 0, '中等(40-59)': 0, '待提高(<40)': 0}
    for score in scores:
        if score >= 80:
            distribution['优秀(80-100)'] += 1
        elif score >= 60:
            distribution['良好(60-79)'] += 1
        elif score >= 40:
            distribution['中等(40-59)'] += 1
        else:
            distribution['待提高(<40)'] += 1
    
    # Category stats
    all_records = ScoreRecord.query.all()
    category_stats = {}
    for record in all_records:
        cat = record.category or record.score_type
        if cat not in category_stats:
            category_stats[cat] = {'total': 0, 'count': 0}
        category_stats[cat]['total'] += record.score
        category_stats[cat]['count'] += 1
    
    for cat in category_stats:
        category_stats[cat]['avg'] = round(category_stats[cat]['total'] / category_stats[cat]['count'], 1)
    
    return jsonify({
        'avg_score': round(sum(scores) / len(scores), 1),
        'max_score': max(scores),
        'min_score': min(scores),
        'score_distribution': distribution,
        'category_stats': category_stats
    })

# AI Analysis API
@app.route('/api/ai/analyze', methods=['POST'])
def ai_analyze():
    client = get_ai_client()
    if not client:
        return jsonify({'error': '请配置 AI API 密钥', 'suggestion': get_mock_suggestion()}), 200
    
    data = request.json
    student_id = data.get('student_id')
    analysis_type = data.get('type', 'general')  # general, improvement, plan
    
    if student_id:
        student = Student.query.get_or_404(student_id)
        records = ScoreRecord.query.filter_by(student_id=student_id)\
            .order_by(ScoreRecord.created_at.desc()).limit(20).all()
        
        context = f"学生姓名: {student.name}\n总分: {student.total_score}\n最近记录:\n"
        for r in records:
            context += f"- {r.created_at.strftime('%m-%d')}: {r.score_type} {r.category} +{r.score}分 ({r.reason})\n"
    else:
        students = Student.query.all()
        context = "全班数据:\n"
        for s in students:
            context += f"- {s.name}: {s.total_score}分\n"
    
    prompts = {
        'general': f"作为班主任，请根据以下量化管理数据，给出综合分析和评价：\n{context}",
        'improvement': f"作为班主任，请分析以下学生的不足之处，并给出具体改进建议：\n{context}",
        'plan': f"作为班主任，请为以下学生制定一个提升计划，包括具体可执行的步骤：\n{context}"
    }
    
    try:
        response = client.chat.completions.create(
            model=AI_MODEL,
            messages=[
                {"role": "system", "content": "你是一位经验丰富的班主任，擅长学生管理和教育。请用简洁专业的语言进行分析。"},
                {"role": "user", "content": prompts.get(analysis_type, prompts['general'])}
            ],
            max_tokens=800,
            temperature=0.7
        )
        analysis = response.choices[0].message.content
        return jsonify({'analysis': analysis, 'type': analysis_type})
    except Exception as e:
        return jsonify({'error': str(e), 'suggestion': get_mock_suggestion()}), 200

def get_mock_suggestion():
    """Mock suggestion when AI is not available"""
    return """
📊 **班级量化管理分析报告**

**整体情况：**
- 建议关注排名靠后的同学，了解具体原因
- 鼓励表现优秀的同学继续保持

**改进方向：**
1. 建立学习互助小组，让优秀学生帮助后进生
2. 定期举办班级活动，增强集体荣誉感
3. 设立阶段性目标，让学生有明确的努力方向

**具体措施：**
- 每周评选"进步之星"
- 建立家校沟通机制
- 开展个性化辅导

*注：这是模拟建议，请配置 AI API 密钥获取智能分析*
    """

# Export API
@app.route('/api/export/excel', methods=['GET'])
def export_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    
    wb = Workbook()
    ws = wb.active
    ws.title = "班级量化统计"
    
    # Header styling
    header_font = Font(bold=True, size=12)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font_white = Font(bold=True, size=12, color="FFFFFF")
    
    # Headers
    headers = ['排名', '姓名', '学号', '性别', '小组', '总分', '备注']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
    
    # Data
    students = Student.query.all()
    students_sorted = sorted(students, key=lambda x: x.total_score, reverse=True)
    
    for row, student in enumerate(students_sorted, 2):
        ws.cell(row=row, column=1, value=row-1)
        ws.cell(row=row, column=2, value=student.name)
        ws.cell(row=row, column=3, value=student.student_id)
        ws.cell(row=row, column=4, value=student.gender)
        ws.cell(row=row, column=5, value=student.group)
        ws.cell(row=row, column=6, value=student.total_score)
        ws.cell(row=row, column=7, value=student.notes)
    
    # Auto width
    for col in ws.columns:
        max_length = 0
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col[0].column_letter].width = max_length + 2
    
    # Save
    filepath = os.path.join(app.instance_path, '班级量化统计.xlsx')
    os.makedirs(app.instance_path, exist_ok=True)
    wb.save(filepath)
    
    return send_file(filepath, as_attachment=True, download_name='班级量化统计.xlsx')

@app.route('/api/export/scores/<int:student_id>', methods=['GET'])
def export_student_scores(student_id):
    from openpyxl import Workbook
    
    student = Student.query.get_or_404(student_id)
    records = ScoreRecord.query.filter_by(student_id=student_id)\
        .order_by(ScoreRecord.created_at.desc()).all()
    
    wb = Workbook()
    ws = wb.active
    ws.title = f"{student.name}评分记录"
    
    # Headers
    headers = ['日期', '类型', '类别', '分数', '原因', '操作人']
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    
    for row, record in enumerate(records, 2):
        ws.cell(row=row, column=1, value=record.created_at.strftime('%Y-%m-%d %H:%M'))
        ws.cell(row=row, column=2, value=record.score_type)
        ws.cell(row=row, column=3, value=record.category)
        ws.cell(row=row, column=4, value=record.score)
        ws.cell(row=row, column=5, value=record.reason)
        ws.cell(row=row, column=6, value=record.operator)
    
    filepath = os.path.join(app.instance_path, f'{student.name}_评分记录.xlsx')
    os.makedirs(app.instance_path, exist_ok=True)
    wb.save(filepath)
    
    return send_file(filepath, as_attachment=True, download_name=f'{student.name}_评分记录.xlsx')

if __name__ == '__main__':
    with app.app_context():
        db.create_all()
    port = int(os.environ.get('PORT', 5000))
    app.run(debug=False, host='0.0.0.0', port=port)
